from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path


try:
    import openpyxl
except Exception:  # pragma: no cover - dependency may be unavailable in some test modes.
    openpyxl = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - dependency may be unavailable in some test modes.
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - dependency may be unavailable in some test modes.
    PdfReader = None


SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".eml", ".xlsx", ".pdf", ".docx"}
FOLDER_CATEGORY_MAP = {
    "photo_intelligence_exports": "photo_intelligence",
    "iqreseller_exports": "iqreseller",
    "email_exports": "email",
    "server_reports": "server_report",
    "misc_docs": "misc",
}


class ParserUnavailableError(RuntimeError):
    pass


@dataclass(slots=True)
class ParsedSegment:
    label: str
    text: str
    metadata: dict[str, object] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedDocument:
    text: str
    source_type: str
    segments: list[ParsedSegment]
    metadata: dict[str, object] = field(default_factory=dict)


def is_supported_file(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def _joined_document(source_type: str, segments: list[ParsedSegment], metadata: dict[str, object] | None = None) -> ParsedDocument:
    text = "\n\n".join(segment.text.strip() for segment in segments if segment.text.strip())
    return ParsedDocument(text=text, source_type=source_type, segments=segments, metadata=metadata or {})


def detect_source_category(path: Path, content_hint: str = "", source_type: str = "") -> str:
    lower_parts = [part.lower() for part in path.parts]
    for folder_name, category in FOLDER_CATEGORY_MAP.items():
        if folder_name in lower_parts:
            return category
    lower_name = path.name.lower()
    hint = f"{lower_name}\n{content_hint[:8000].lower()}"
    if "photo_intelligence" in hint or "printed_serial" in hint or "barcode_serial" in hint or "photo_names" in hint:
        return "photo_intelligence"
    if "server_report" in hint or "pass_with_warnings" in hint or "job=" in hint or "missing report" in hint:
        return "server_report"
    if source_type == "eml" or lower_name.endswith(".eml") or ("subject:" in hint and ("from:" in hint or "to:" in hint)):
        return "email"
    if "iqreseller" in hint:
        return "iqreseller"
    if "demo" in lower_parts:
        return "demo"
    return "misc"


def _require_dependency(value: object, package_name: str) -> None:
    if value is None:
        raise ParserUnavailableError(f"{package_name} is required for this parser but is not installed.")


def _paragraph_segments(text: str, *, label_prefix: str) -> list[ParsedSegment]:
    blocks = [block.strip() for block in text.replace("\r\n", "\n").split("\n\n") if block.strip()]
    if not blocks:
        blocks = [line.strip() for line in text.splitlines() if line.strip()]
    if not blocks:
        return []
    return [ParsedSegment(label=f"{label_prefix} {index:04d}", text=block) for index, block in enumerate(blocks, start=1)]


def _grouped_line_segments(text: str, *, lines_per_segment: int, label_prefix: str) -> list[ParsedSegment]:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").splitlines() if line.strip()]
    segments: list[ParsedSegment] = []
    for start in range(0, len(lines), lines_per_segment):
        block = lines[start : start + lines_per_segment]
        if not block:
            continue
        first_line = start + 1
        last_line = start + len(block)
        label = f"{label_prefix} lines {first_line:04d}-{last_line:04d}"
        segments.append(ParsedSegment(label=label, text="\n".join(block)))
    return segments


def _parse_csv(path: Path) -> ParsedDocument:
    segments: list[ParsedSegment] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample or "a,b\n1,2\n")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(handle, dialect=dialect)
        rows = list(reader)
    if not rows:
        return _joined_document("csv", [ParsedSegment(label="row 0001", text=f"CSV file: {path.name} is empty.")])
    header = rows[0]
    data_rows = rows[1:] if header else rows
    for row_index, row in enumerate(data_rows, start=1):
        pairs = []
        for column_index, value in enumerate(row):
            column_name = header[column_index] if column_index < len(header) and header[column_index] else f"column_{column_index + 1}"
            pairs.append(f"{column_name}={value}")
        text = f"source_file={path.name} | row_number={row_index} | " + " | ".join(pairs)
        segments.append(ParsedSegment(label=f"row {row_index:04d}", text=text, metadata={"row_number": row_index}))
    if not segments:
        segments.append(ParsedSegment(label="row 0001", text=f"CSV file: {path.name} has headers only."))
    return _joined_document("csv", segments)


def _flatten_json(prefix: str, value: object, lines: list[str]) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            _flatten_json(child_prefix, child, lines)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            child_prefix = f"{prefix}[{index}]"
            _flatten_json(child_prefix, child, lines)
    else:
        lines.append(f"{prefix}={value}")


def _parse_json(path: Path) -> ParsedDocument:
    content = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    lines: list[str] = []
    _flatten_json("", content, lines)
    segments: list[ParsedSegment] = []
    for index in range(0, len(lines), 20):
        block = lines[index : index + 20]
        label = f"json block {index // 20 + 1:04d}"
        segments.append(ParsedSegment(label=label, text="\n".join(block)))
    return _joined_document("json", segments or [ParsedSegment(label="json block 0001", text=f"JSON file: {path.name} is empty.")])


def _parse_textual(path: Path, source_type: str) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    if source_type == "log":
        segments = _grouped_line_segments(text, lines_per_segment=20, label_prefix="log")
    else:
        segments = _paragraph_segments(text, label_prefix=source_type)
    if not segments:
        segments = [ParsedSegment(label=f"{source_type} 0001", text=f"{path.name} is empty.")]
    return _joined_document(source_type, segments)


def _parse_eml(path: Path) -> ParsedDocument:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    headers = [
        f"From: {message.get('From', '')}",
        f"To: {message.get('To', '')}",
        f"Subject: {message.get('Subject', '')}",
        f"Date: {message.get('Date', '')}",
    ]
    if message.is_multipart():
        body_parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                body_parts.append(part.get_content())
        body = "\n\n".join(body_parts).strip()
    else:
        body = message.get_content().strip()
    header_text = "\n".join(headers).strip()
    body_segments = _paragraph_segments(body, label_prefix="email body") if body else []
    if not body_segments:
        body_segments = [ParsedSegment(label="email body 0001", text="[No body content found]")]
    full_text = header_text + "\n\n" + "\n\n".join(segment.text for segment in body_segments)
    return ParsedDocument(
        text=full_text.strip(),
        source_type="eml",
        segments=body_segments,
        metadata={"context_prefix": header_text},
    )


def _parse_xlsx(path: Path) -> ParsedDocument:
    _require_dependency(openpyxl, "openpyxl")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    segments: list[ParsedSegment] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_row = [str(value).strip() if value is not None else "" for value in rows[0]]
        for row_number, row in enumerate(rows[1:], start=2):
            if not any(value not in (None, "") for value in row):
                continue
            pairs = []
            for column_index, value in enumerate(row):
                column_name = header_row[column_index] if column_index < len(header_row) and header_row[column_index] else f"column_{column_index + 1}"
                pairs.append(f"{column_name}={value}")
            text = f"source_file={path.name} | sheet_name={sheet.title} | row_number={row_number} | " + " | ".join(pairs)
            segments.append(
                ParsedSegment(
                    label=f"{sheet.title} row {row_number:04d}",
                    text=text,
                    metadata={"sheet_name": sheet.title, "row_number": row_number},
                )
            )
    if not segments:
        segments.append(ParsedSegment(label="sheet 0001", text=f"Workbook {path.name} contains no populated rows."))
    return _joined_document("xlsx", segments)


def _parse_pdf(path: Path) -> ParsedDocument:
    _require_dependency(PdfReader, "pypdf")
    reader = PdfReader(str(path))
    segments: list[ParsedSegment] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            text = "[No extractable text found on this page]"
        segments.append(
            ParsedSegment(
                label=f"page {page_index:04d}",
                text=text,
                metadata={"page_number": page_index},
            )
        )
    if not segments:
        segments.append(ParsedSegment(label="page 0001", text=f"PDF file {path.name} contains no pages."))
    return _joined_document("pdf", segments)


def _parse_docx(path: Path) -> ParsedDocument:
    _require_dependency(DocxDocument, "python-docx")
    document = DocxDocument(str(path))
    segments: list[ParsedSegment] = []
    paragraph_index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_index += 1
        segments.append(ParsedSegment(label=f"paragraph {paragraph_index:04d}", text=text))
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if not cells:
                continue
            segments.append(
                ParsedSegment(
                    label=f"table {table_index:02d} row {row_index:04d}",
                    text=" | ".join(cells),
                    metadata={"table_index": table_index, "row_number": row_index},
                )
            )
    if not segments:
        segments.append(ParsedSegment(label="paragraph 0001", text=f"DOCX file {path.name} contains no extractable text."))
    return _joined_document("docx", segments)


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".json":
        return _parse_json(path)
    if suffix in {".txt", ".md", ".log"}:
        return _parse_textual(path, suffix.lstrip("."))
    if suffix == ".eml":
        return _parse_eml(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")
